"""Write the gap assessment workbook."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ecc.assess import Assessment, compliance, domain_scores
from ecc.load import Dataset
from ecc.mapping import controls_for
from ecc.models import Status
from ecc.risk import risk_rows, roadmap

SHEETS = ("Summary", "Controls", "Technical evidence", "Risk register", "Roadmap")
STATUS_COLOURS = {
    Status.IMPLEMENTED: "C6EFCE",
    Status.PARTIAL: "FFEB9C",
    Status.NOT_IMPLEMENTED: "FFC7CE",
    Status.NOT_APPLICABLE: "D9D9D9",
}
RESULT_COLOURS = {
    "passed": "C6EFCE",
    "failed": "FFC7CE",
    "not applicable": "D9D9D9",
    "not scanned": "D9D9D9",
}


def fill(colour: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


def table(
    sheet: Worksheet,
    header: Sequence[str],
    rows: Iterable[Sequence[Any]],
    widths: Sequence[int],
    *,
    autofilter: bool = False,
) -> None:
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if autofilter:
        sheet.auto_filter.ref = sheet.dimensions


def score(value: float | None) -> float | str:
    return "n/a" if value is None else value


def change(before: float | None, after: float | None) -> float | str:
    if before is None or after is None:
        return "n/a"
    return round(after - before, 1)


def label(status: Status | None, empty: str = "") -> str:
    return status.value if status is not None else empty


def write_summary(
    sheet: Worksheet, dataset: Dataset, after: Assessment, before: Assessment | None
) -> None:
    titles = {c.domain: f"{c.domain} {c.domain_title}" for c in dataset.catalog}
    after_domains = domain_scores(after)
    before_domains = domain_scores(before) if before is not None else {}
    overall_before = compliance(before.results) if before is not None else None
    areas = [("Overall", compliance(after.results), overall_before)]
    areas += [(titles[d], after_domains[d], before_domains.get(d)) for d in after_domains]
    if before is None:
        header = ["Area", f"{after.snapshot.label} %"]
        rows = [[area, score(now)] for area, now, _ in areas]
    else:
        header = ["Area", f"{before.snapshot.label} %", f"{after.snapshot.label} %", "Change"]
        rows = [[area, score(old), score(now), change(old, now)] for area, now, old in areas]
    table(sheet, header, rows, [48, 14, 14, 10])
    for row in sheet.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0"
    sheet.append([])
    sheet.append(["Status", "Controls"])
    for column in (1, 2):
        sheet.cell(row=sheet.max_row, column=column).font = Font(bold=True)
    for status in Status:
        sheet.append([status.value, sum(1 for r in after.results if r.status is status)])
        sheet.cell(row=sheet.max_row, column=1).fill = fill(STATUS_COLOURS[status])
    sheet.append(["No rating", sum(1 for r in after.results if r.status is None)])
    sheet.append(["Total", len(after.results)])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    snapshot = after.snapshot
    sheet.append([])
    sheet.append(
        [
            "As of",
            f"{snapshot.taken_at:%d %B %Y}, snapshot '{snapshot.label}', "
            f"workstation {snapshot.host}",
        ]
    )
    scoring = (
        "Compliance = (Implemented + 0.5 x Partially implemented) / rated applicable controls "
        "(Not applicable and unrated controls are left out). "
        "n/a = no applicable rated controls."
    )
    if before is not None:
        scoring += " Change is in percentage points."
    sheet.append(["Scoring", scoring])
    sheet.append(
        [
            "Roadmap",
            "Buckets count days from the snapshot date; overdue items fall in '30 days'.",
        ]
    )


def write_controls(sheet: Worksheet, after: Assessment) -> None:
    rows = [
        [
            r.control.id,
            f"{r.control.domain} {r.control.domain_title}",
            f"{r.control.subdomain} {r.control.subdomain_title}",
            r.control.text,
            label(r.status, "No rating"),
            label(r.technical),
            label(r.documentary),
            r.summary,
        ]
        for r in after.results
    ]
    header = ["ID", "Domain", "Subdomain", "Control", "Status", "Technical", "Documentary",
              "Evidence"]
    table(sheet, header, rows, [9, 24, 30, 60, 22, 22, 22, 60], autofilter=True)
    for row_number, result in enumerate(after.results, start=2):
        if result.status is not None:
            sheet.cell(row=row_number, column=5).fill = fill(STATUS_COLOURS[result.status])


def write_checks(sheet: Worksheet, dataset: Dataset, after: Assessment) -> None:
    results = {check.id: check.result for check in after.snapshot.checks}
    rows = [
        [
            check.id,
            check.cis,
            check.title,
            results.get(check.id, "not scanned"),
            ", ".join(controls_for(check, dataset.mapping)),
        ]
        for check in dataset.policy.checks
    ]
    table(
        sheet,
        ["Check ID", "CIS", "Title", "Result", "ECC controls"],
        rows,
        [10, 12, 80, 16, 18],
        autofilter=True,
    )
    for row_number, row in enumerate(rows, start=2):
        sheet.cell(row=row_number, column=4).fill = fill(RESULT_COLOURS[row[3]])


def write_risks(register: Worksheet, plan: Worksheet, dataset: Dataset, after: Assessment) -> None:
    rows = risk_rows(dataset.risks, after)
    table(
        register,
        ["ID", "Controls", "Risk", "Likelihood", "Impact", "Score", "Rating", "Treatment",
         "Owner", "Due", "State"],
        [
            [
                row.risk.id,
                ", ".join(row.risk.controls),
                row.risk.statement,
                row.risk.likelihood,
                row.risk.impact,
                row.score,
                row.rating,
                row.risk.treatment,
                row.risk.owner,
                row.risk.due,
                "Closed" if row.closed else "Open",
            ]
            for row in rows
        ],
        [8, 14, 60, 11, 9, 8, 10, 11, 20, 12, 9],
        autofilter=True,
    )
    table(
        plan,
        ["When", "Risk ID", "Controls", "Risk", "Score", "Owner", "Due"],
        [
            [
                row.bucket,
                row.risk.id,
                ", ".join(row.risk.controls),
                row.risk.statement,
                row.score,
                row.risk.owner,
                row.risk.due,
            ]
            for row in roadmap(rows)
        ],
        [10, 9, 14, 60, 8, 20, 12],
    )


def write_workbook(
    path: Path, dataset: Dataset, after: Assessment, before: Assessment | None = None
) -> None:
    book = Workbook()
    summary = book.active
    summary.title = SHEETS[0]
    sheets = [summary, *(book.create_sheet(title) for title in SHEETS[1:])]
    write_summary(sheets[0], dataset, after, before)
    write_controls(sheets[1], after)
    write_checks(sheets[2], dataset, after)
    write_risks(sheets[3], sheets[4], dataset, after)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
